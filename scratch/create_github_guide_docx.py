import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_callout(doc, text_list, title="💡 CATATAN PENTING"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0F9FF")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Border left blue
    tcPr = cell._element.get_or_add_tcPr()
    tcBorders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="0284C7"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(tcBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"{title}\n")
    run_t.bold = True
    run_t.font.name = "Segoe UI"
    run_t.font.size = Pt(10.5)
    run_t.font.color.rgb = RGBColor(2, 132, 199)
    
    for item in text_list:
        p_item = cell.add_paragraph()
        p_item.paragraph_format.space_before = Pt(0)
        p_item.paragraph_format.space_after = Pt(2)
        r = p_item.add_run(f"• {item}")
        r.font.name = "Segoe UI"
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(51, 65, 85)
    
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def build_document():
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Document Header Title
    p_header = doc.add_paragraph()
    p_header.paragraph_format.space_before = Pt(0)
    p_header.paragraph_format.space_after = Pt(4)
    r_sub = p_header.add_run("PANDUAN LENGKAP PENGELOLAAN REPOSITORY GITHUB")
    r_sub.font.name = "Segoe UI"
    r_sub.font.size = Pt(10)
    r_sub.font.bold = True
    r_sub.font.color.rgb = RGBColor(2, 132, 199)
    
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(12)
    r_main = p_title.add_run("Ringkasan Perintah Git, Keterangan, Fungsi, & Workflow Kolaborasi")
    r_main.font.name = "Segoe UI"
    r_main.font.size = Pt(20)
    r_main.font.bold = True
    r_main.font.color.rgb = RGBColor(15, 23, 42)
    
    # Meta box
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(16)
    r_meta = p_meta.add_run("Repository: inawawi/admincyberstore  |  Branch Utamanya: main, dev_cyberstoreimw (IMW), dev_cyberstoremlh (MLH)")
    r_meta.font.name = "Segoe UI"
    r_meta.font.size = Pt(9.5)
    r_meta.font.italic = True
    r_meta.font.color.rgb = RGBColor(100, 116, 139)

    add_callout(doc, [
        "Dokumen ini dibuat khusus untuk memandu kolaborasi pengembangan aplikasi UBSI Cyber Store.",
        "Branch IMW (dev_cyberstoreimw) digunakan oleh Bapak Imam Nawawi.",
        "Branch MLH (dev_cyberstoremlh) digunakan oleh Mas Muhammad Luthfi Hamid.",
        "Setiap fitur/perubahan di-commit di branch masing-masing sebelum digabungkan (merge)."
    ], title="📌 RINGKASAN REPOSITORY COLLABORATION")

    # SECTION 1: PERINTAH DASAR GIT (BASIC WORKFLOW)
    h1 = doc.add_heading(level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. Perintah Dasar Git (Basic Workflow)")
    r_h1.font.name = "Segoe UI"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(15, 23, 42)
    
    cmd_basics = [
        ("git status", "Memeriksa Status Repository", "Menampilkan status direktori kerja saat ini, termasuk file mana yang telah diubah (modified), baru ditambahkan (untracked), atau siap disimpan (staged)."),
        ("git add .", "Menandai Semua Perubahan (Stage)", "Memasukkan seluruh file yang baru dibuat atau diubah ke dalam staging area agar siap dimasukkan ke dalam commit berikutnya. Tanda titik (.) artinya semua file di folder project."),
        ("git add <nama_file>", "Menandai File Spesifik", "Hanya memasukkan file tertentu yang ditentukan ke staging area. Contoh: git add components/resource-client.tsx."),
        ("git commit -m \"pesan\"", "Menyimpan Perubahan (Commit)", "Menyimpan snapshot perubahan dari staging area ke riwayat Git lokal lengkap dengan pesan deskriptif. Contoh: git commit -m \"Update tampilan 3 kolom\"."),
        ("git push origin <nama_branch>", "Mengunggah Kode ke GitHub", "Mengirimkan commit yang ada di komputer lokal Anda ke server GitHub pada branch tujuan. Contoh: git push origin dev_cyberstoreimw."),
        ("git pull origin <nama_branch>", "Mengunduh & Menggabungkan Update", "Mengambil commit terbaru dari branch di server GitHub dan langsung menggabungkannya ke branch lokal Anda saat ini. Contoh: git pull origin dev_cyberstoreimw."),
        ("git fetch origin", "Mengunduh Metadata Tanpa Merge", "Mengecek dan mengunduh informasi terbaru dari GitHub tanpa mengubah file kodingan Anda di komputer lokal.")
    ]
    
    for cmd, label, desc in cmd_basics:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(2)
        r_cmd = p_c.add_run(f"💻 {cmd}")
        r_cmd.bold = True
        r_cmd.font.name = "Consolas"
        r_cmd.font.size = Pt(10.5)
        r_cmd.font.color.rgb = RGBColor(2, 132, 199)
        
        r_lbl = p_c.add_run(f" — {label}")
        r_lbl.bold = True
        r_lbl.font.name = "Segoe UI"
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(15, 23, 42)
        
        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(6)
        p_d.paragraph_format.left_indent = Inches(0.25)
        r_d = p_d.add_run(desc)
        r_d.font.name = "Segoe UI"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # SECTION 2: PENGELOLAAN BRANCH (BRANCH MANAGEMENT)
    h2 = doc.add_heading(level=1)
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. Perintah Pengelolaan Branch (Branch Management)")
    r_h2.font.name = "Segoe UI"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(15, 23, 42)

    cmd_branches = [
        ("git branch", "Melihat Daftar Branch Lokal", "Menampilkan seluruh branch yang ada di komputer Anda. Branch yang sedang aktif ditandai dengan bintang (*) dan warna hijau."),
        ("git branch -a", "Melihat Semua Branch (Lokal & Remote)", "Menampilkan seluruh branch baik yang ada di komputer lokal maupun yang ada di server GitHub (remote/origin)."),
        ("git checkout <nama_branch>", "Pindah ke Branch Lain", "Mengubah branch aktif Anda ke branch lain yang ditentukan. Contoh: git checkout dev_cyberstoreimw untuk berpindah ke branch milik IMW."),
        ("git checkout -b <branch_baru>", "Membuat & Pindah Branch Baru", "Membuat branch baru dari posisi saat ini sekaligus langsung berpindah ke branch baru tersebut. Contoh: git checkout -b fitur_baru."),
        ("git merge <nama_branch>", "Menggabungkan Kode Branch", "Menggabungkan commit dari branch lain ke branch yang sedang aktif saat ini. Contoh jika di main: git merge dev_cyberstoreimw."),
        ("git branch -d <nama_branch>", "Menghapus Branch Lokal", "Menghapus branch lokal yang sudah tidak digunakan (setelah di-merge).")
    ]

    for cmd, label, desc in cmd_branches:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(2)
        r_cmd = p_c.add_run(f"🌿 {cmd}")
        r_cmd.bold = True
        r_cmd.font.name = "Consolas"
        r_cmd.font.size = Pt(10.5)
        r_cmd.font.color.rgb = RGBColor(16, 185, 129)

        r_lbl = p_c.add_run(f" — {label}")
        r_lbl.bold = True
        r_lbl.font.name = "Segoe UI"
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(15, 23, 42)

        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(6)
        p_d.paragraph_format.left_indent = Inches(0.25)
        r_d = p_d.add_run(desc)
        r_d.font.name = "Segoe UI"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # SECTION 3: WORKFLOW KOLABORASI REPOSITORY ADMINCYBERSTORE
    h3 = doc.add_heading(level=1)
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. Panduan Langkah Praktis Kolaborasi AdminCyberStore")
    r_h3.font.name = "Segoe UI"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = RGBColor(15, 23, 42)

    # Subsection A
    p_sa = doc.add_paragraph()
    p_sa.paragraph_format.space_before = Pt(4)
    p_sa.paragraph_format.space_after = Pt(4)
    r_sa = p_sa.add_run("A. Alur Rutin Pengelolaan Hasil Kerja IMW (dev_cyberstoreimw):")
    r_sa.bold = True
    r_sa.font.name = "Segoe UI"
    r_sa.font.size = Pt(11)
    r_sa.font.color.rgb = RGBColor(2, 132, 199)

    steps_imw = [
        "Pastikan berada di branch IMW: git checkout dev_cyberstoreimw",
        "Periksa perubahan file: git status",
        "Tandai semua perubahan: git add .",
        "Simpan commit dengan pesan jelas: git commit -m \"Update fitur ...\"",
        "Kirim ke GitHub: git push origin dev_cyberstoreimw"
    ]
    for idx, st in enumerate(steps_imw, 1):
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Inches(0.2)
        p_st.paragraph_format.space_after = Pt(2)
        r_num = p_st.add_run(f"{idx}. ")
        r_num.bold = True
        r_num.font.name = "Segoe UI"
        r_txt = p_st.add_run(st)
        r_txt.font.name = "Consolas" if "git" in st else "Segoe UI"
        r_txt.font.size = Pt(9.5)

    # Subsection B
    p_sb = doc.add_paragraph()
    p_sb.paragraph_format.space_before = Pt(10)
    p_sb.paragraph_format.space_after = Pt(4)
    r_sb = p_sb.add_run("B. Cara Mengambil Update Terbaru dari Rekan (MLH / Luthfi):")
    r_sb.bold = True
    r_sb.font.name = "Segoe UI"
    r_sb.font.size = Pt(11)
    r_sb.font.color.rgb = RGBColor(2, 132, 199)

    steps_mlh = [
        "Jika Anda ingin mengambil update dari Luthfi ke branch Anda:",
        "  1. git checkout dev_cyberstoreimw",
        "  2. git pull origin dev_cyberstoremlh",
        "Jika Luthfi ingin mengambil update terbaru dari Anda ke branch-nya:",
        "  1. git checkout dev_cyberstoremlh",
        "  2. git pull origin dev_cyberstoreimw"
    ]
    for st in steps_mlh:
        p_st = doc.add_paragraph()
        p_st.paragraph_format.left_indent = Inches(0.2)
        p_st.paragraph_format.space_after = Pt(2)
        r_txt = p_st.add_run(st)
        r_txt.font.name = "Consolas" if "git" in st else "Segoe UI"
        r_txt.font.size = Pt(9.5)

    # SECTION 4: TROUBLESHOOTING & UNDO
    h4 = doc.add_heading(level=1)
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r_h4 = h4.add_run("4. Perintah Diagnosa & Pembatalan (Troubleshooting)")
    r_h4.font.name = "Segoe UI"
    r_h4.font.size = Pt(14)
    r_h4.font.bold = True
    r_h4.font.color.rgb = RGBColor(15, 23, 42)

    cmd_trouble = [
        ("git log --oneline -n 5", "Melihat 5 Riwayat Commit Terakhir", "Menampilkan daftar riwayat commit terakhir secara singkat 1 baris per commit."),
        ("git diff", "Melihat Perincian Perubahan Kode", "Menampilkan perbedaan baris kode yang diubah sebelum di-commit."),
        ("git stash", "Menyimpan Perubahan Sementara", "Menyimpan sementara kodingan yang belum di-commit ke dalam kantong simpanan (stash) tanpa perlu di-commit."),
        ("git stash pop", "Mengeluarkan Simpanan Sementara", "Mengembalikan kembali kodingan yang sebelumnya disimpan dengan git stash."),
        ("git restore <nama_file>", "Membatalkan Edit File (Belum Staged)", "Membatalkan perubahan pada file dan mengembalikannya seperti versi commit terakhir.")
    ]

    for cmd, label, desc in cmd_trouble:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.space_before = Pt(4)
        p_c.paragraph_format.space_after = Pt(2)
        r_cmd = p_c.add_run(f"🛠️ {cmd}")
        r_cmd.bold = True
        r_cmd.font.name = "Consolas"
        r_cmd.font.size = Pt(10.5)
        r_cmd.font.color.rgb = RGBColor(217, 70, 239)

        r_lbl = p_c.add_run(f" — {label}")
        r_lbl.bold = True
        r_lbl.font.name = "Segoe UI"
        r_lbl.font.size = Pt(10)
        r_lbl.font.color.rgb = RGBColor(15, 23, 42)

        p_d = doc.add_paragraph()
        p_d.paragraph_format.space_before = Pt(0)
        p_d.paragraph_format.space_after = Pt(6)
        p_d.paragraph_format.left_indent = Inches(0.25)
        r_d = p_d.add_run(desc)
        r_d.font.name = "Segoe UI"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = RGBColor(51, 65, 85)

    # SECTION 5: TABEL CHEAT SHEET RINGKASAN
    h5 = doc.add_heading(level=1)
    h5.paragraph_format.space_before = Pt(18)
    h5.paragraph_format.space_after = Pt(8)
    r_h5 = h5.add_run("5. Tabel Ringkasan Perintah Git (Cheat Sheet)")
    r_h5.font.name = "Segoe UI"
    r_h5.font.size = Pt(14)
    r_h5.font.bold = True
    r_h5.font.color.rgb = RGBColor(15, 23, 42)

    table_data = [
        ("git status", "Informasi", "Cek status file yang diubah / belum ditandai", "git status"),
        ("git checkout <branch>", "Branch", "Berpindah ke branch lain yang dituju", "git checkout dev_cyberstoreimw"),
        ("git add .", "Staging", "Menandai seluruh file perubahan untuk commit", "git add ."),
        ("git commit -m \"msg\"", "Commit", "Menyimpan perubahan lokal dengan pesan", "git commit -m \"Update UI\""),
        ("git push origin <branch>", "Remote", "Mengunggah commit ke server GitHub", "git push origin dev_cyberstoreimw"),
        ("git pull origin <branch>", "Remote", "Mengunduh & gabung update dari GitHub", "git pull origin dev_cyberstoremlh"),
        ("git branch -a", "Branch", "Melihat daftar seluruh branch lokal & remote", "git branch -a"),
        ("git merge <branch>", "Branch", "Menggabungkan isi branch ke branch aktif", "git merge dev_cyberstoremlh"),
        ("git log --oneline", "History", "Melihat riwayat commit secara ringkas", "git log --oneline -n 5"),
        ("git restore <file>", "Undo", "Membatalkan edit pada file sebelum commit", "git restore app/globals.css")
    ]

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header Row
    hdr_cells = table.rows[0].cells
    headers = ["PERINTAH GIT", "KATEGORI", "FUNGSI UTAMA", "CONTOH PENGGUNAAN"]
    widths = [Inches(1.8), Inches(1.0), Inches(2.3), Inches(1.9)]
    
    for idx, h_text in enumerate(headers):
        hdr_cells[idx].width = widths[idx]
        set_cell_background(hdr_cells[idx], "0F172A")
        set_cell_margins(hdr_cells[idx], top=120, bottom=120, left=100, right=100)
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h_text)
        r.bold = True
        r.font.name = "Segoe UI"
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for row_idx, data_row in enumerate(table_data):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(data_row):
            row_cells[col_idx].width = widths[col_idx]
            set_cell_background(row_cells[col_idx], bg_color)
            set_cell_margins(row_cells[col_idx], top=90, bottom=90, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            r = p.add_run(cell_value)
            r.font.size = Pt(8.5)
            if col_idx in (0, 3):
                r.font.name = "Consolas"
                r.font.color.rgb = RGBColor(2, 132, 199)
            else:
                r.font.name = "Segoe UI"
                r.font.color.rgb = RGBColor(51, 65, 85)

    # Save Document
    out_path = r"d:\BTI\UBSI CYBER STORE\cyber_store_nextjs\docs\Ringkasan_Perintah_Git_GitHub.docx"
    import os
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print(f"SUCCESS: {out_path}")

if __name__ == "__main__":
    build_document()
